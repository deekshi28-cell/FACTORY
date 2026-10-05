import pymupdf  
import os
import pytesseract
from PIL import Image
import io
import docx
from docx.oxml.ns import qn
import openpyxl
import zipfile
import logging

logger = logging.getLogger("readers")



def read_pdf(filepath):
    """
    Reads a PDF and returns a list of pages, each with page number and text.
    Works for text-based PDFs. Scanned PDFs will return empty/near-empty text.
    """
    try:
        doc = pymupdf.open(filepath)
    except Exception as e:
        logger.error("Failed to open PDF '%s': %s", filepath, e, exc_info=True)
        raise

    pages = []
    try:
        for page_num in range(len(doc)):
            try:
                page = doc[page_num]
                text = page.get_text()
            except Exception as e:
                logger.warning("Error reading page %d of '%s': %s", page_num + 1, filepath, e)
                text = ""
            pages.append({
                "page_number": page_num + 1,
                "text": text,
                "source_file": os.path.basename(filepath)
            })
    finally:
        doc.close()
    return pages


def read_pdf_with_ocr(filepath, ocr_lang="jpn+eng"):
    """
    Reads a PDF page by page. If a page has little/no extractable text,
    falls back to OCR on that page's rendered image.
    """
    try:
        doc = pymupdf.open(filepath)
    except Exception as e:
        logger.error("Failed to open PDF with OCR '%s': %s", filepath, e, exc_info=True)
        raise

    pages = []
    try:
        for page_num in range(len(doc)):
            page = doc[page_num]
            try:
                text = page.get_text()
            except Exception as e:
                logger.warning("Failed to extract text from page %d of '%s': %s", page_num + 1, filepath, e)
                text = ""
            used_ocr = False

            if len(text.strip()) < 20:
                try:
                    pix = page.get_pixmap(dpi=300)
                    img = Image.open(io.BytesIO(pix.tobytes("png")))
                    text = pytesseract.image_to_string(img, lang=ocr_lang)
                    used_ocr = True
                except Exception as e:
                    logger.error("OCR failed for '%s' page %d: %s", filepath, page_num + 1, e)

            pages.append({
                "page_number": page_num + 1,
                "text": text,
                "source_file": os.path.basename(filepath),
                "used_ocr": used_ocr
            })
    finally:
        doc.close()
    return pages


def _read_table_grid(table):
    
    num_cols = len(table.columns)
    grid = []
    for row in table.rows:
        grid_row = [None] * num_cols
        col_idx = 0
        tr = row._tr
        for tc in tr.findall(qn('w:tc')):
            tcPr = tc.find(qn('w:tcPr'))
            vmerge = None
            gridspan = 1
            if tcPr is not None:
                vm = tcPr.find(qn('w:vMerge'))
                gs = tcPr.find(qn('w:gridSpan'))
                if vm is not None:
                    vmerge = vm.get(qn('w:val'), 'continue')
                if gs is not None:
                    gridspan = int(gs.get(qn('w:val')))

            text = ''.join(t.text for t in tc.iter(qn('w:t'))).strip()

            while col_idx < num_cols and grid_row[col_idx] is not None:
                col_idx += 1

            if vmerge == 'continue':
                text = ''

            for i in range(gridspan):
                if col_idx + i < num_cols:
                    grid_row[col_idx + i] = text if i == 0 else ''
            col_idx += gridspan

        grid.append([c if c is not None else '' for c in grid_row])
    return grid


def read_docx(filepath):
    try:
        doc = docx.Document(filepath)
    except Exception as e:
        logger.error("Failed to open DOCX '%s': %s", filepath, e, exc_info=True)
        raise

    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)

    tables = []
    for table_idx, table in enumerate(doc.tables):
        try:
            rows = _read_table_grid(table)
        except Exception as e:
            logger.warning("Skipping unreadable table %d in '%s': %s", table_idx, filepath, e)
            rows = []
        tables.append({
            "table_index": table_idx,
            "rows": rows
        })

    return {
        "source_file": os.path.basename(filepath),
        "paragraphs": paragraphs,
        "tables": tables
    }


def read_xlsx(filepath):
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
    except Exception as e:
        logger.error("Failed to open XLSX '%s': %s", filepath, e, exc_info=True)
        raise

    sheets = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                rows.append(list(row))
        sheets.append({
            "sheet_name": sheet_name,
            "rows": rows
        })
    return {
        "source_file": os.path.basename(filepath),
        "sheets": sheets
    }


def extract_images_from_pdf(filepath, output_folder, min_size=250, max_size=1800):
    os.makedirs(output_folder, exist_ok=True)
    try:
        doc = pymupdf.open(filepath)
    except Exception as e:
        logger.error("Failed to open PDF for image extraction '%s': %s", filepath, e, exc_info=True)
        return []

    base_name = os.path.splitext(os.path.basename(filepath))[0]
    images = []

    try:
        for page_num in range(len(doc)):
            try:
                page = doc[page_num]
                image_list = page.get_images(full=True)
            except Exception as e:
                logger.warning("Failed to read images on page %d of '%s': %s", page_num + 1, filepath, e)
                continue

            for img_index, img in enumerate(image_list):
                try:
                    xref = img[0]
                    base_image = doc.extract_image(xref)
                    width = base_image["width"]
                    height = base_image["height"]

                    if width < min_size or height < min_size:
                        continue
                    if width > max_size or height > max_size:
                        continue

                    ext = base_image["ext"]
                    img_filename = f"{base_name}_p{page_num+1}_img{img_index}.{ext}"
                    img_path = os.path.join(output_folder, img_filename)

                    with open(img_path, "wb") as f:
                        f.write(base_image["image"])

                    images.append({
                        "image_path": img_path,
                        "page_number": page_num + 1,
                        "source_file": os.path.basename(filepath)
                    })
                except Exception as e:
                    logger.warning(
                        "Skipping unreadable image %d on page %d of '%s': %s",
                        img_index, page_num + 1, filepath, e
                    )
    finally:
        doc.close()

    return images


def extract_images_from_docx(filepath, output_folder, min_size=150):
    
    os.makedirs(output_folder, exist_ok=True)
    base_name = os.path.splitext(os.path.basename(filepath))[0]
    images = []
    try:
        with zipfile.ZipFile(filepath, 'r') as z:
            img_idx = 0
            for name in z.namelist():
                if 'media/' in name and not name.endswith('/'):
                    data = z.read(name)
                    try:
                        img = Image.open(io.BytesIO(data))
                        w, h = img.size
                        if w < min_size or h < min_size:
                            continue
                        ext = os.path.splitext(name)[1].lstrip('.').lower()
                        if not ext:
                            ext = "png"
                        img_filename = f"{base_name}_docx_img{img_idx}.{ext}"
                        img_path = os.path.join(output_folder, img_filename)
                        with open(img_path, "wb") as f:
                            f.write(data)
                        images.append({
                            "image_path": img_path,
                            "page_number": 0,
                            "source_file": os.path.basename(filepath)
                        })
                        img_idx += 1
                    except Exception as e:
                        logger.warning("Skipping unreadable image '%s' in DOCX '%s': %s", name, filepath, e)
    except Exception as e:
        logger.error("Failed to extract images from DOCX '%s': %s", filepath, e, exc_info=True)
    return images


def extract_images_from_xlsx(filepath, output_folder, min_size=150):
   
    os.makedirs(output_folder, exist_ok=True)
    base_name = os.path.splitext(os.path.basename(filepath))[0]
    images = []
    try:
        with zipfile.ZipFile(filepath, 'r') as z:
            img_idx = 0
            for name in z.namelist():
                if 'media/' in name and not name.endswith('/'):
                    data = z.read(name)
                    try:
                        img = Image.open(io.BytesIO(data))
                        w, h = img.size
                        if w < min_size or h < min_size:
                            continue
                        ext = os.path.splitext(name)[1].lstrip('.').lower()
                        if not ext:
                            ext = "png"
                        img_filename = f"{base_name}_xlsx_img{img_idx}.{ext}"
                        img_path = os.path.join(output_folder, img_filename)
                        with open(img_path, "wb") as f:
                            f.write(data)
                        images.append({
                            "image_path": img_path,
                            "page_number": 0,
                            "source_file": os.path.basename(filepath)
                        })
                        img_idx += 1
                    except Exception as e:
                        logger.warning("Skipping unreadable image '%s' in XLSX '%s': %s", name, filepath, e)
    except Exception as e:
        logger.error("Failed to extract images from XLSX '%s': %s", filepath, e, exc_info=True)
    return images


if __name__ == "__main__":
    
    print("=== Verifying merged-cell fix on 6.docx ===")
    test_file = r"D:\FactoryKA\documents\6.docx"
    result = read_docx(test_file)

    for idx in [30, 31]:
        print(f"\n--- Table {idx} ---")
        for row in result["tables"][idx]["rows"]:
            print(row)